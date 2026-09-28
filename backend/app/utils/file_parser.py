"""
Best-effort text extraction for uploaded learning materials (PDF, DOCX, PPTX, TXT).
Extraction failures never raise - they fall back to an empty string so the
upload itself always succeeds and the quiz generator can fall back to the
topic-based question bank.
"""
from pathlib import Path


def extract_text(path: Path, file_type: str) -> str:
    file_type = (file_type or "").lower().lstrip(".")
    try:
        if file_type == "txt":
            return path.read_text(encoding="utf-8", errors="ignore")

        if file_type == "pdf":
            try:
                from pypdf import PdfReader
            except ImportError:
                from PyPDF2 import PdfReader  # type: ignore
            reader = PdfReader(str(path))
            return "\n".join((page.extract_text() or "") for page in reader.pages)

        if file_type == "docx":
            import docx  # python-docx
            doc = docx.Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs)

        if file_type == "pptx":
            from pptx import Presentation
            prs = Presentation(str(path))
            chunks = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        chunks.append(shape.text)
            return "\n".join(chunks)
    except Exception:
        return ""
    return ""
